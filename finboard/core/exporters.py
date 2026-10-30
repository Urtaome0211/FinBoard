"""Excel / PDF report exports (PDF rendering lives in the UI layer)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .dates import month_label, today_iso
from .models import TransactionFilter
from .money import percentage
from .repository import Repository


def _money(cents: int) -> float:
    return round(cents / 100, 2)


def export_excel(
    repo: Repository,
    path: str | Path,
    month: str | None = None,
) -> Path:
    target = Path(path)
    net = repo.net_worth()
    stats = repo.month_stats(month)
    breakdown = repo.category_breakdown(month, "expense")
    transactions = repo.transactions(
        TransactionFilter(month=month, limit=100_000)
    )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "概览"
    summary["A1"] = "FinBoard 财务报表"
    summary["A1"].font = Font(size=16, bold=True)
    summary["A2"] = "统计月份"
    summary["B2"] = month_label(month) if month else "全部"
    summary["A3"] = "导出时间"
    summary["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    summary["A5"] = "资产"
    summary["B5"] = _money(net.assets)
    summary["A6"] = "负债"
    summary["B6"] = _money(net.liabilities)
    summary["A7"] = "净资产"
    summary["B7"] = _money(net.net)
    summary["A8"] = "本月收入"
    summary["B8"] = _money(stats.income)
    summary["A9"] = "本月支出"
    summary["B9"] = _money(stats.expense)
    summary["A10"] = "本月结余"
    summary["B10"] = _money(stats.net)
    for cell in ("B5", "B6", "B7", "B8", "B9", "B10"):
        summary[cell].number_format = '#,##0.00'
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 20

    category_sheet = workbook.create_sheet("支出分类")
    category_sheet.append(["分类", "金额", "占比"])
    total = sum(row.amount for row in breakdown)
    for row in breakdown:
        category_sheet.append(
            [row.name, _money(row.amount), percentage(row.amount, total) / 100]
        )
    category_sheet["C1"].number_format = "0.0%"
    for row in range(2, category_sheet.max_row + 1):
        category_sheet.cell(row=row, column=2).number_format = '#,##0.00'
    category_sheet.column_dimensions["A"].width = 22
    category_sheet.column_dimensions["B"].width = 16
    category_sheet.column_dimensions["C"].width = 12

    sheet = workbook.create_sheet("收支明细")
    headers = ["日期", "类型", "金额", "币种", "分类", "账户", "转入账户", "到账金额", "手续费", "对方", "备注", "标签"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3D6DF2")
        cell.alignment = Alignment(horizontal="center")
    from .models import KIND_LABELS

    for transaction in transactions:
        sheet.append(
            [
                transaction.date,
                KIND_LABELS.get(transaction.kind, transaction.kind),
                _money(transaction.amount),
                transaction.currency,
                transaction.category_name,
                transaction.account_name,
                transaction.to_account_name,
                _money(transaction.to_amount) if transaction.to_amount else None,
                _money(transaction.fee) if transaction.fee else None,
                transaction.payee,
                transaction.note,
                ", ".join(tag.name for tag in transaction.tags),
            ]
        )
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=3).number_format = '#,##0.00'
        if sheet.cell(row=row, column=8).value is not None:
            sheet.cell(row=row, column=8).number_format = '#,##0.00'
        if sheet.cell(row=row, column=9).value is not None:
            sheet.cell(row=row, column=9).number_format = '#,##0.00'
    widths = [12, 10, 14, 8, 16, 18, 18, 14, 12, 18, 26, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for ws in (summary, category_sheet, sheet):
        ws.sheet_view.showGridLines = False
    workbook.save(target)
    return target


def report_filename(prefix: str = "FinBoard-报表", extension: str = "xlsx") -> str:
    return f"{prefix}-{today_iso()}.{extension}"
